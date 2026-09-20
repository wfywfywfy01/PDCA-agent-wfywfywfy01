# -*- coding: utf-8 -*-
"""以清理后的 872 条为主键，回原始名单匹配，输出原格式的已清理文件。

主键：国家／地区 + 商场 + 店铺 + 服务类型 + 电话（五字段组合，源内唯一）
输出：完整保留源文件的 sheet 结构 / 列 / 版式，只把被剔除的行滤掉，
      并在「剩余概览」页追加清理说明。
"""
import sys, os, shutil
from collections import Counter
import openpyxl
from copy import copy

sys.stdout = open(r"D:\经销商PDCA\.tmp_xlsx\match.log", "w", encoding="utf-8")

SRC = r"C:\Users\frank\Desktop\02_本轮未分配剩余名单.xlsx"
CLEAN = r"C:\Users\frank\Desktop\高端商场维修线索_2026-09-16\01_剩余名单_已清理品牌直营.xlsx"
OUT = r"C:\Users\frank\Desktop\02_本轮未分配剩余名单_已清理品牌直营.xlsx"


def norm(v):
    if v is None: return ""
    s = " ".join(str(v).split())
    return "" if s.lower() in ("none", "nan", "-", "n/a", "na", "null") else s


def read(path, sheet):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = [[norm(c) for c in r] for r in wb[sheet].iter_rows(values_only=True)]
    wb.close()
    return rows


# ---------- 1. 建立主键集合（来自清理后的 872 条） ----------
crow = read(CLEAN, "剩余名单（已清理）")
CH = crow[0]; CB = [r for r in crow[1:] if any(r)]
CI = {h: j for j, h in enumerate(CH)}
PKF = ("国家／地区", "商场", "店铺", "服务类型", "电话")
def ckey(r): return tuple(r[CI[f]] if CI.get(f, -1) < len(r) else "" for f in PKF)
clean_keys = Counter_ = None
from collections import Counter
clean_keys = Counter(ckey(r) for r in CB)
print(f"主键来源：{CLEAN}")
print(f"  行数 {len(CB)}  唯一主键 {len(clean_keys)}  （重复键 {len(CB)-len(clean_keys)}）")

# ---------- 2. 源名单 ----------
srow = read(SRC, "剩余名单")
SH = srow[0]; SB = [r for r in srow[1:] if any(r)]
SI = {h: j for j, h in enumerate(SH)}
def skey(r): return tuple(r[SI[f]] if SI.get(f, -1) < len(r) else "" for f in PKF)
src_keys = Counter(skey(r) for r in SB)
print(f"\n源名单：{SRC}")
print(f"  行数 {len(SB)}  唯一主键 {len(src_keys)}  （重复键 {len(SB)-len(src_keys)}）")

matched_in_src = [k for k in clean_keys if k in src_keys]
print(f"\n主键匹配：清理版 {len(clean_keys)} 个键 → 在源名单命中 {len(matched_in_src)} 个")
print(f"  清理版有、源名单没有 = {len(clean_keys)-len(matched_in_src)}")
ghost = [k for k in clean_keys if k not in src_keys]
for k in ghost[:8]:
    print(f"     {k[2][:40]} | {k[1][:30]} | {k[4][:20]}")

keep_set = set(clean_keys)
dropped = [r for r in SB if skey(r) not in keep_set]
kept = [r for r in SB if skey(r) in keep_set]
print(f"\n源名单 {len(SB)} = 保留 {len(kept)} + 剔除 {len(dropped)}")

# ---------- 3. 输出（复制源文件，删掉被剔除的行，保留版式） ----------
shutil.copy2(SRC, OUT)
wb = openpyxl.load_workbook(OUT)
ws = wb["剩余名单"]

# 找表头行
hdr_row = 1
for i in range(1, 6):
    if "国家／地区" in [norm(ws.cell(row=i, column=c).value) for c in range(1, ws.max_column + 1)]:
        hdr_row = i; break
col_of = {norm(ws.cell(row=hdr_row, column=c).value): c for c in range(1, ws.max_column + 1)}
print(f"\n输出表头行 = {hdr_row}，列数 = {ws.max_column}")

pk_cols = [col_of[f] for f in PKF]
data_rows = ws.max_row
to_del = []
for r in range(hdr_row + 1, data_rows + 1):
    k = tuple(norm(ws.cell(row=r, column=c).value) for c in pk_cols)
    if k not in keep_set:
        to_del.append(r)
print(f"待删除行数 = {len(to_del)}")

for r in reversed(to_del):
    ws.delete_rows(r, 1)
print(f"删除后数据行 = {ws.max_row - hdr_row}")

# ---------- 4. 剩余概览页追加清理说明 ----------
ov = wb["剩余概览"]
row = ov.max_row + 2
def put(ws, r, c, v, bold=False):
    cell = ws.cell(row=r, column=c, value=v)
    if bold: cell.font = copy(ws.cell(row=1, column=1).font); cell.font = cell.font.copy(bold=True)

ov.cell(row=row, column=1, value="本轮清理（品牌直营剔除）").font = openpyxl.styles.Font(bold=True)
row += 1
for lab, val in [
    ("清理规则", "剔除品牌直营／品牌专卖店、单品牌授权服务商（只服务自家产品，无经销可能）"),
    ("清理前", f"{len(SB)} 条"),
    ("清理后", f"{len(kept)} 条（本表）"),
    ("已剔除", f"{len(dropped)} 条"),
    ("匹配方式", "以「国家／地区+商场+店铺+服务类型+电话」五字段主键与清理后的可用名单逐条匹配"),
    ("保留内容", "原 29 列一字未改，未新增列"),
]:
    ov.cell(row=row, column=1, value=lab)
    ov.cell(row=row, column=2, value=val)
    row += 1
wb.save(OUT)
print(f"\n已保存：{OUT}")

# ---------- 5. 复核 ----------
wb2 = openpyxl.load_workbook(OUT, read_only=True, data_only=True)
w2 = wb2["剩余名单"]
rows2 = [[norm(c) for c in r] for r in w2.iter_rows(values_only=True)]
H2 = rows2[hdr_row - 1]
B2 = [r for r in rows2[hdr_row:] if any(r)]
I2 = {h: j for j, h in enumerate(H2)}
keys2 = set(tuple(r[I2[f]] for f in PKF) for r in B2)
wb2.close()
print("\n" + "=" * 70); print("复核"); print("=" * 70)
print(f"  输出文件数据行 = {len(B2)}")
print(f"  输出文件唯一主键 = {len(keys2)}")
print(f"  与清理版主键完全一致 = {keys2 == keep_set}")
print(f"  sheet 列表 = {openpyxl.load_workbook(OUT, read_only=True).sheetnames}")
print(f"  列数 = {len([h for h in H2 if h])}")
