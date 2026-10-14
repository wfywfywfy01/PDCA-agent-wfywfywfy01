"""Extract bounded text from a meeting upload without retaining the source file."""
from __future__ import annotations

import io
import re
from pathlib import PurePosixPath
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader
from pypdf.errors import PdfReadError


MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARS = 1800
_MAX_UNZIPPED_BYTES = 12 * 1024 * 1024
_EXTENSIONS = {".txt", ".md", ".csv", ".pdf", ".docx", ".pptx", ".xlsx"}


def _checked_zip(data: bytes) -> ZipFile:
    archive = ZipFile(io.BytesIO(data))
    entries = archive.infolist()
    if len(entries) > 500 or sum(part.file_size for part in entries) > _MAX_UNZIPPED_BYTES:
        archive.close()
        raise HTTPException(413, "文档展开后过大，请上传精简版")
    return archive


def _ooxml_text(data: bytes, extension: str) -> str:
    with _checked_zip(data) as archive:
        names = archive.namelist()
        if extension == ".docx":
            parts = ["word/document.xml"] if "word/document.xml" in names else []
            tag = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
        else:
            parts = sorted((name for name in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
                           key=lambda name: int(re.search(r"\d+", name).group()))
            tag = "{http://schemas.openxmlformats.org/drawingml/2006/main}t"
        if not parts:
            raise ValueError("文档缺少可读取的正文")
        lines = []
        for part in parts[:40]:
            root = ElementTree.fromstring(archive.read(part))
            lines.extend(node.text or "" for node in root.iter(tag))
            if sum(map(len, lines)) > MAX_TEXT_CHARS:
                break
        return "\n".join(lines)


def _spreadsheet_text(data: bytes) -> str:
    with _checked_zip(data):
        pass
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    try:
        lines = []
        for sheet in book.worksheets[:5]:
            lines.append(f"工作表：{sheet.title}")
            for row in sheet.iter_rows(max_row=500, max_col=20, values_only=True):
                line = " | ".join(str(value) for value in row if value is not None)
                if line:
                    lines.append(line)
                if sum(map(len, lines)) > MAX_TEXT_CHARS:
                    return "\n".join(lines)
        return "\n".join(lines)
    finally:
        book.close()


def _pdf_text(data: bytes) -> str:
    if not data.startswith(b"%PDF-"):
        raise ValueError("PDF 文件头无效")
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ValueError("加密 PDF 暂不支持")
    lines = []
    for page in reader.pages[:40]:
        lines.append(page.extract_text() or "")
        if sum(map(len, lines)) > MAX_TEXT_CHARS:
            break
    return "\n".join(lines)


def extract_material(filename: str, data: bytes) -> dict:
    if not data:
        raise HTTPException(422, "文件为空")
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "单个资料不能超过 5 MB")
    name = re.split(r"[/\\]", filename or "")[-1].strip()
    if not name or len(name) > 120 or any(ord(char) < 32 for char in name):
        raise HTTPException(422, "文件名无效或超过 120 字")
    extension = PurePosixPath(name).suffix.lower()
    if extension not in _EXTENSIONS:
        raise HTTPException(422, "支持 TXT、MD、CSV、PDF、DOCX、PPTX、XLSX")
    try:
        if extension in {".txt", ".md", ".csv"}:
            if b"\x00" in data:
                raise ValueError("文本文件含二进制内容")
            try:
                raw = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                raw = data.decode("gb18030")
        elif extension in {".docx", ".pptx"}:
            raw = _ooxml_text(data, extension)
        elif extension == ".xlsx":
            raw = _spreadsheet_text(data)
        else:
            raw = _pdf_text(data)
    except (BadZipFile, ElementTree.ParseError, UnicodeDecodeError, InvalidFileException, ValueError,
            KeyError, OSError, PdfReadError) as exc:
        raise HTTPException(422, "资料无法读取，请检查文件格式或另存为文本") from exc
    raw = re.sub(r"[ \t]+", " ", raw).strip()
    raw = re.sub(r"\n{3,}", "\n\n", raw)
    if not raw:
        raise HTTPException(422, "未提取到文字；扫描图片请先转成可复制文字的文件")
    return {"name": name, "text": raw[:MAX_TEXT_CHARS], "truncated": len(raw) > MAX_TEXT_CHARS}
